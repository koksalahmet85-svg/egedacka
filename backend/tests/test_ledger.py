"""Tahsilat defteri: parser, sezon analizi ve API. Gerçek defter verisi KULLANILMAZ —
gerçek dosyanın düzensizliklerini taklit eden sentetik bir çalışma kitabı üretilir."""
import io
from datetime import date, datetime

import openpyxl
import pytest

from app.services.ledger_import import classify, parse_money, parse_sheet_month, parse_workbook, player_key
from app.services.season_analysis import Row, compare

HDR = ["TARİH", "OYUNCU AD SOYAD", "GRUP", "DÖNEM ", "TUTAR", "AİDAT", "FORMA", "AÇIKLAMA"]


def build_workbook() -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    # Gerçek defterdeki gibi: bazı sayfalar B sütunundan, bazıları A'dan başlıyor.
    sheets = {
        ("TEMMUZ 25", 1): [
            [datetime(2025, 7, 1), "TEST OYUNCU A", "SPOR OKULU", "TEMMUZ", "₺4.300,00", "₺4.300,00", "-", "POS"],
            ["21.07.2025", "TEST OYUNCU B", "ALTYAPI", "SEZONLUK PEŞİN ÖDEME", "₺72.000,00", "₺72.000,00", "-", "POS"],
            ["21.07.2025", "TEST OYUNCU C", "ALTYAPI", "TEMMUZ", 4700, 4700, "-", "FİNANS"],
        ],
        ("AĞUSTOS 25", 1): [
            ["01.08.205", "TEST OYUNCU A", "SPOR OKULU", "AĞUSTOS", 4300, 4300, "-", "GARANTİ"],
            ["19.08.2025", "TEST OYUNCU D", "ALTYAPI", "SEZONLUK PEŞİN ÖDEME", 170000, 72000, "-", "POS"],
            ["20.08.2025", "TEST OYUNCU C", "ALTYAPI", "TEMMUZ+AĞUSTOS", 9400, 9400, "-", "MAİL ORDER"],
            ["21.08.2025", "TEST OYUNCU E", "ALT YAPI", "İSTANBUL ", 20000, 20000, "-", "NAKİT"],
            ["26.08.2025", "TEST OYUNCU F", "ALTYAPI", "AĞUSTOS", 5200, "₺5,200.00", "-", "MAİL ORDER"],
            ["27.08.2025", "TEST OYUNCU G", "ALTYAPI", "ÖDEME İPTAL EDİLECEK", 85000, 85000, "-", "SANAL POS İPTAL"],
        ],
        ("TEMMUZ 26", 0): [
            [datetime(2026, 7, 2), "TEST OYUNCU A", "SPOR OKULU", "HAZİRAN", "CEVAD", 7450, "-", "GARANTİ"],
            ["29.07.2026", "TEST OYUNCU C", "ALTYAPI", "SEZONLUK PEŞİN", 119000, 119000, "-", "POS"],
        ],
        ("AĞUSTOS 26", 0): [
            ["02.08.2026", "TEST OYUNCU A", "SPOR OKULU", "AĞUSTOS + TEKSTİL", 14250, 8250, 6000, "POS"],
            ["03.08.2026", "TEST OYUNCU H", "ALTYAPI", "AĞUSTOS", 9000, 9000, "-", "GARANTİ"],
            ["11 08. 2026", "TEST OYUNCU F", "ALTYAPI", "SEZONLUK 1.TAKSİT", 60000, 60000, 0, "MAİL ORDER"],
        ],
    }
    for (title, col0), data in sheets.items():
        ws = wb.create_sheet(title)
        ws.append([None] * col0 + HDR)
        for r in data:
            ws.append([None] * col0 + r)
    ws = wb.create_sheet("ÖZET")  # ay adı olmayan sayfa atlanmalı
    ws.append(["toplam"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def parsed():
    return parse_workbook(build_workbook())


def test_helpers():
    assert parse_sheet_month("AĞUSTOS 25") == date(2025, 8, 1)
    assert parse_sheet_month("ÖZET") is None
    assert str(parse_money("₺4.700,00")) == "4700.00"
    assert str(parse_money("₺5,200.00")) == "5200.00"
    assert parse_money("-") is None and parse_money("CEVAD") is None
    assert player_key("Çağan  Çınar Ar") == player_key("ÇAĞAN ÇINAR AR") == "CAGAN CINAR AR"
    kind, months, _, _ = classify("ARALIK+OCAK", date(2026, 1, 1))
    assert kind == "MULTI_MONTH" and months == [date(2025, 12, 1), date(2026, 1, 1)]
    assert classify("SEZONLUK 2. TAKSİT", date(2025, 10, 1))[0] == "SEASON_INSTALLMENT"
    assert classify("SEZONLUK PEŞİN BAŞARI İNDİRİMİ", date(2026, 7, 1))[2] is True
    assert classify("EYLÜL YARIM AİDAT", date(2025, 9, 1))[0] == "PARTIAL"
    assert classify("YEŞİL TRİBÜN", date(2026, 10, 1))[0] == "EVENT"
    assert classify("TEKSTİL (GECEN SENE)", date(2026, 9, 1))[0] == "TEXTILE"


def test_parse_workbook(parsed):
    assert parsed.sheets == ["TEMMUZ 25", "AĞUSTOS 25", "TEMMUZ 26", "AĞUSTOS 26"]
    by = {(r.sheet, r.player_name): r for r in parsed.records}
    assert len(parsed.records) == 14
    # Yıl hatalı tarih sayfa yılıyla düzeltilir ve işaretlenir
    a = by[("AĞUSTOS 25", "TEST OYUNCU A")]
    assert a.paid_on == date(2025, 8, 1) and a.paid_on_inferred
    # Tekstilsiz tutar uyuşmazlığı: küçük olan alınır (gelir şişmez)
    d = by[("AĞUSTOS 25", "TEST OYUNCU D")]
    assert float(d.amount) == 72000 and d.kind == "SEASON_UPFRONT"
    # Çok aylık etiket, önceki ayı doğru yıla oturtur
    assert by[("AĞUSTOS 25", "TEST OYUNCU C")].covered_months == [date(2025, 7, 1), date(2025, 8, 1)]
    assert by[("AĞUSTOS 25", "TEST OYUNCU E")].kind == "EVENT"
    assert by[("AĞUSTOS 25", "TEST OYUNCU E")].group == "ALTYAPI"  # "ALT YAPI"
    assert by[("AĞUSTOS 25", "TEST OYUNCU G")].cancelled
    # Metin tutar → AİDAT+FORMA'dan türetilir
    assert float(by[("TEMMUZ 26", "TEST OYUNCU A")].amount) == 7450
    t = by[("AĞUSTOS 26", "TEST OYUNCU A")]
    assert (float(t.dues_amount), float(t.textile_amount)) == (8250, 6000)
    assert by[("AĞUSTOS 26", "TEST OYUNCU F")].paid_on == date(2026, 8, 11)  # "11 08. 2026"
    assert by[("TEMMUZ 25", "TEST OYUNCU C")].channel == "FINANSBANK"
    issues = {x.issue for x in parsed.anomalies}
    assert {"tarih", "tutar_uyumsuz", "tutar", "iptal", "sayfa_atlandi"} <= issues


def _all_rows(parsed):
    return [Row(r.sheet_month, r.paid_on, r.player_key, r.group, r.kind, tuple(r.covered_months), r.scholarship,
                float(r.amount), float(r.dues_amount), float(r.textile_amount), r.channel, r.cancelled)
            for r in parsed.records]


def test_season_rows_and_compare(parsed):
    from app.services.season_analysis import season_rows

    rows = _all_rows(parsed)
    b0, b1, t0, t1 = date(2025, 8, 1), date(2026, 7, 31), date(2026, 8, 1), date(2027, 7, 31)

    # Kasa: yalnız sezon aralığındaki sayfalar. Temmuz 25 (önceki sezon) dışarıda, Temmuz 26 peşini içeride.
    cb, _ = season_rows(rows, b0, b1, "cash")
    assert sum(r.amount for r in cb if not r.cancelled) == 4300 + 72000 + 9400 + 20000 + 5200 + 7450 + 119000

    # Ekonomik: Temmuz 25 erken kayıt peşini 25-26'ya girer, Temmuz 25 aylık aidatları girmez;
    # Temmuz 26 peşini 26-27'ye taşınır.
    eb, b_carried = season_rows(rows, b0, b1)
    et, t_carried = season_rows(rows, t0, t1)
    assert b_carried == {date(2025, 7, 1)} and t_carried == {date(2026, 7, 1)}
    assert sum(r.amount for r in eb if not r.cancelled) == 72000 + 4300 + 72000 + 9400 + 20000 + 5200 + 7450
    assert sum(r.amount for r in et) == 119000 + 14250 + 9000 + 60000

    eco = compare("A", eb, "B", et, "economic", target_carried=t_carried)
    assert eco["target"]["season_plan"]["upfront_amount"] == 119000
    w = eco["aligned"]["window"]
    assert w["months"] == ["2026-07", "2026-08"] and w["season_only_months"] == ["2026-07"]
    # Hedefin son ödemesi 11 Ağustos → bazda Temmuz 25 peşini (72.000) + 1 Ağustos aidatı kıyasa girer.
    assert w["last_month_cutoff_day"] == 11
    assert eco["aligned"]["season_plan_amount"]["base"] == 72000
    assert eco["aligned"]["total"]["base"] == 72000 + 4300
    assert eco["retention"]["recent_retained"] >= 1  # C Temmuz 26'da peşin ödedi


async def test_import_and_compare_api(client, admin, conn):
    for name, s, e in (("2025-2026", "2025-08-01", "2026-07-31"), ("2026-2027", "2026-08-01", "2027-07-31")):
        await conn.execute("DELETE FROM seasons WHERE name=$1 AND NOT EXISTS (SELECT 1 FROM teams t WHERE t.season_id=seasons.id)", name)
        r = await client.put("/seasons", json={"name": name, "start_date": s, "end_date": e}, headers=admin)
        assert r.status_code == 200, r.text

    content = build_workbook()
    files = {"file": ("defter.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r1 = await client.post("/ledger/import", files=files, headers=admin)
    assert r1.status_code == 201, r1.text
    assert r1.json()["rows"] == 14
    r2 = await client.post("/ledger/import", files=files, headers=admin)  # aynı dosya tekrar
    assert r2.json()["replaced_rows"] == 14
    assert await conn.fetchval("SELECT COUNT(*) FROM ledger_entries") == 14

    bad = await client.post("/ledger/import", files={"file": ("x.csv", b"a,b", "text/csv")}, headers=admin)
    assert bad.status_code == 422

    c = await client.get("/ledger/compare?base=2025-2026&target=2026-2027", headers=admin)
    assert c.status_code == 200, c.text
    assert c.json()["mode"] == "economic"
    assert (await client.get("/ledger/compare?base=yok&target=2026-2027", headers=admin)).status_code == 404

    nr = (await client.get("/ledger/not-returned?base=2025-2026&target=2026-2027", headers=admin)).json()
    keys = {x["player_key"] for x in nr}
    assert "TEST OYUNCU C" not in keys  # sezonluk peşin ödedi
    assert "TEST OYUNCU A" not in keys  # hedef sezonda ödedi

    anomalies = (await client.get("/ledger/anomalies", headers=admin)).json()
    assert any(a["issue"] == "tutar_uyumsuz" for a in anomalies)

    # İptal satırı nakit akışına girmez; tekstil ayrı kategori
    cats = {r["category"]: float(r["s"]) for r in await conn.fetch(
        "SELECT category, SUM(amount) AS s FROM v_cash_movements WHERE category LIKE 'LEDGER%' GROUP BY 1")}
    assert cats["LEDGER_TEXTILE"] == 6000 and cats["LEDGER_EVENT"] == 20000


async def test_ledger_requires_finance_role(client, login):
    coach = await login("coach1@t")
    assert (await client.get("/ledger/compare?base=2025-2026&target=2026-2027", headers=coach)).status_code == 403


def test_upfront_vs_monthly_12_month_coverage():
    from app.services.season_analysis import coverage_months, upfront_vs_monthly

    cov = coverage_months(date(2027, 7, 31))
    assert cov[0] == date(2026, 8, 1) and cov[-1] == date(2027, 7, 1) and len(cov) == 12

    def monthly(m, price, i):
        return Row(m, m, f"P{i}", "ALTYAPI", "MONTHLY", (m,), False, price, price, 0.0, "POS")

    rows = [monthly(date(2026, 8, 1), 9000.0, i) for i in range(6)]
    rows += [monthly(date(2026, 9, 1), 9900.0, i) for i in range(6)]
    rows += [Row(date(2026, 7, 1), date(2026, 7, 1), f"U{i}", "ALTYAPI", "SEASON_UPFRONT", (), False,
                 119000.0, 119000.0, 0.0, "POS") for i in range(5)]
    rows += [Row(date(2026, 7, 1), date(2026, 7, 1), "X", "ALTYAPI", "SEASON_UPFRONT", (), False,
                 200000.0, 200000.0, 0.0, "POS")]  # tekil fiyat noktası: gürültü, raporlanmaz
    (v,) = upfront_vs_monthly(rows, cov)
    # Ağu 9.000 + Eyl 9.900 + kalan 10 ay son fiyatla (zam yok varsayımı)
    assert v["monthly_path"] == 9000 + 11 * 9900 and v["projected_months"] == 10
    (p,) = v["price_points"]
    assert p["price"] == 119000 and p["discount_vs_monthly"] == pytest.approx(1 - 119000 / 117900)
    # 119.000 - (9.000 + 9.900) = 100.100 → kalan 10 ayda %1,1 zam başabaş
    assert p["breakeven_hike"] == pytest.approx(100100 / 99000 - 1)

    # Ocak 2027'den itibaren %20 zam senaryosu: Eki–Ara zamsız, Oca–Tem (7 ay) zamlı
    (h,) = upfront_vs_monthly(rows, cov, date(2027, 1, 1), 0.20)
    assert h["monthly_path"] == pytest.approx(9000 + 9900 + 3 * 9900 + 7 * 9900 * 1.2)
    assert h["price_points"][0]["breakeven_hike"] == pytest.approx((119000 - 9000 - 4 * 9900) / (7 * 9900) - 1)


def test_plan_tiers_early_bird_and_deadline():
    from app.services.season_analysis import coverage_months, plan_tiers

    cov = coverage_months(date(2027, 7, 31))
    rows = [Row(date(2026, 8, 1), date(2026, 8, 1), f"M{i}", "ALTYAPI", "MONTHLY", (date(2026, 8, 1),), False,
                9900.0, 9900.0, 0.0, "POS") for i in range(5)]

    def plan(key, paid, amount, kind="SEASON_UPFRONT"):
        return Row(date(paid.year, paid.month, 1), paid, key, "ALTYAPI", kind, (), False, amount, amount, 0.0, "POS")

    rows += [plan("E1", date(2026, 7, 20), 107000), plan("E2", date(2026, 8, 5), 107000),  # E2 son tarih sonrası
             plan("S1", date(2026, 7, 25), 119000),
             plan("T1", date(2026, 7, 28), 59500, "SEASON_INSTALLMENT"),  # taksitli standart plan
             plan("T1", date(2026, 9, 28), 59500, "SEASON_INSTALLMENT"),
             plan("X", date(2026, 8, 20), 130000)]
    tiers = [{"group": "ALTYAPI", "tier": "EARLY", "label": "Erken kayıt", "price": 107000.0,
              "valid_until": date(2026, 7, 31)},
             {"group": "ALTYAPI", "tier": "STANDARD", "label": None, "price": 119000.0, "valid_until": None}]
    (g,) = plan_tiers(rows, tiers, cov)
    early, std = g["tiers"]
    assert g["monthly_path"] == 12 * 9900
    assert (early["players"], early["after_deadline"]) == (2, 1)
    assert (std["players"], std["upfront"], std["installment"]) == (2, 1, 1)
    assert early["discount_vs_monthly"] == pytest.approx(1 - 107000 / 118800)
    assert g["other_price_players"] == 1 and g["other_price_amount"] == 130000
    assert g["in_progress_players"] == 0


async def test_plan_prices_api(client, admin):
    body = [{"group": "ALTYAPI", "tier": "EARLY", "label": "Erken kayıt", "price": 107000},
            {"group": "ALTYAPI", "tier": "STANDARD", "price": 119000}]
    r = await client.put("/seasons/2026-2027/plan-prices", json=body, headers=admin)
    assert r.status_code == 200, r.text
    assert len((await client.get("/seasons/2026-2027/plan-prices", headers=admin)).json()) == 2
    assert (await client.put("/seasons/yok/plan-prices", json=body, headers=admin)).status_code == 404
    c = (await client.get("/ledger/compare?base=2025-2026&target=2026-2027", headers=admin)).json()
    assert c["plan_tiers"]["target"][0]["tiers"][1]["players"] == 1  # sentetik defterde 119.000 peşin


def test_ledger_diff_detects_added_changed_moved_removed():
    from app.services.ledger_diff import diff

    old = parse_workbook(build_workbook())
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook()))
    ws = wb["AĞUSTOS 26"]
    ws.insert_rows(2)  # başlığın altına satır ekle → altındaki satırlar kayar (MOVED)
    ws.cell(2, 1, "05.08.2026"); ws.cell(2, 2, "TEST OYUNCU Z"); ws.cell(2, 3, "ALTYAPI")
    ws.cell(2, 4, "AĞUSTOS"); ws.cell(2, 5, 9000); ws.cell(2, 6, 9000); ws.cell(2, 8, "POS")
    t = wb["TEMMUZ 26"]
    for row in t.iter_rows(min_row=2):
        if row[1].value == "TEST OYUNCU C":
            row[3].value = "SEZONLUK 1. TAKSİT"  # aynı satır, aynı oyuncu, farklı dönem → CHANGED
    buf = io.BytesIO(); wb.save(buf)
    new = parse_workbook(buf.getvalue())
    res = diff(old, new)
    counts = res.counts()
    assert counts.get("ADDED") == 1 and counts.get("CHANGED") == 1
    assert counts.get("MOVED", 0) >= 1 and "REMOVED" not in counts
    assert res.net_delta() == 9000
    added = [c for c in res.changes if c.kind == "ADDED"][0]
    assert added.sheet == "AĞUSTOS 26" and added.new_row == 2

    assert diff(old, old).changes == []
